from pathlib import Path
import pytest
from backend.ingestion import parse_document, local_files
from backend.config import settings


def test_markdown_encoding_and_sources(tmp_path):
    file = tmp_path / "知识.md"
    file.write_text("# Java 基础\n\nString 是不可变字符串。\n" * 200, encoding="utf-8")
    chunks = parse_document(file)
    assert len(chunks) > 1 and all(len(text) <= 1800 for text, _ in chunks)
    assert all("知识.md" in source for _, source in chunks)


def test_docx_tables_are_preserved(tmp_path):
    from docx import Document
    doc = Document(); doc.add_paragraph("Java 知识")
    table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = "volatile"; table.cell(0, 1).text = "可见性"
    path = tmp_path / "sample.docx"; doc.save(path)
    content = "".join(x for x, _ in parse_document(path))
    assert "volatile" in content and "可见性" in content


def test_docx_textboxes_nested_tables_and_alternate_content(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml
    doc = Document()
    paragraph = doc.add_paragraph("Java图示")
    paragraph._p.append(parse_xml('''<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><mc:Choice Requires="w"><w:txbxContent><w:p><w:r><w:t>图中Eden区域</w:t></w:r></w:p></w:txbxContent></mc:Choice><mc:Fallback><w:txbxContent><w:p><w:r><w:t>图中Eden区域</w:t></w:r></w:p></w:txbxContent></mc:Fallback></mc:AlternateContent>'''))
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "外层表格"
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "内层CAS知识"
    path = tmp_path / "drawing.docx"; doc.save(path)
    content = "".join(text for text, _ in parse_document(path))
    assert content.count("图中Eden区域") == 1
    assert "外层表格" in content and "内层CAS知识" in content


def test_empty_pdf_requires_ocr(tmp_path):
    from pypdf import PdfWriter
    writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
    file = tmp_path / "empty.pdf"; writer.write(file)
    with pytest.raises(ValueError, match="OCR"):
        parse_document(file)


def test_local_import_cannot_escape_authorized_root(tmp_path, monkeypatch):
    root = tmp_path / "allowed"; root.mkdir(); (root / "data.md").write_text("知识", encoding="utf-8")
    monkeypatch.setattr(settings(), "import_root", root)
    assert len(local_files(str(root))) == 1
    with pytest.raises(ValueError, match="IMPORT_ROOT"):
        local_files(str(tmp_path))


def test_import_source_migration_preserves_legacy_jobs():
    from sqlalchemy import create_engine, text, inspect
    from backend.migrations import migrate_import_sources
    engine = create_engine("sqlite://")
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE import_jobs (id VARCHAR(32) PRIMARY KEY, status VARCHAR(24), total INTEGER)"))
        conn.execute(text("INSERT INTO import_jobs VALUES ('old', 'failed', 4)"))
    migrate_import_sources(engine); migrate_import_sources(engine)
    assert {"source_paths", "source_kind"}.issubset({c["name"] for c in inspect(engine).get_columns("import_jobs")})
    with engine.connect() as conn:
        row = conn.execute(text("SELECT status, total, source_paths FROM import_jobs WHERE id='old'")).one()
        assert row.status == "failed" and row.total == 4 and row.source_paths is None
    engine.dispose()
